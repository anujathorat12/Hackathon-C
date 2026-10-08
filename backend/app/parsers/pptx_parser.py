import os
from typing import Optional, Dict, Any, List
from pptx import Presentation
from ..shared.schemas.template_models import TemplateGuidanceProfile
from .template_normalizer import normalize_guidance
from .theme_reader import read_theme, palette_from_theme

def is_vibrant_accent(r: int, g: int, b: int) -> bool:
    """Checks if an RGB color is a vibrant accent (not monochrome black/white/gray)."""
    max_c = max(r, g, b)
    min_c = min(r, g, b)
    diff = max_c - min_c
    return diff > 30 and 40 < max_c < 245

def calculate_luminance(r: int, g: int, b: int) -> float:
    """Calculates relative luminance (0..255 range)."""
    return 0.2126 * r + 0.7152 * g + 0.0722 * b

def parse_pptx_template(file_path: str) -> TemplateGuidanceProfile:
    """
    Inspects an uploaded PowerPoint (.pptx) template file.
    Dynamically extracts slide layouts, font families, dark/light theme background, and accent colors.
    """
    if not os.path.exists(file_path):
        return normalize_guidance()

    try:
        prs = Presentation(file_path)
        base_name = os.path.basename(file_path)

        available_layouts: List[str] = []
        heading_font = "Calibri"
        primary_hex = "#1B365D"
        accent_hex = "#6C5CE7"
        background_hex = "#FFFFFF"
        text_hex = "#222222"

        found_backgrounds: List[str] = []
        detected_sections: List[str] = []
        found_accents: List[str] = []
        has_light_text: bool = False

        # 1. Inspect Slide Layouts in Slide Master
        for layout in prs.slide_layouts:
            name = layout.name.strip() if layout.name else "Custom Layout"
            if name not in available_layouts:
                available_layouts.append(name)

        # 2. Inspect Sample Shapes and Text Frames for Fonts, Backgrounds, and Colors
        for slide in prs.slides:
            title_found = False
            # Check slide background fill
            try:
                if hasattr(slide, 'background') and slide.background and slide.background.fill:
                    fill = slide.background.fill
                    if hasattr(fill, 'fore_color') and fill.fore_color and fill.fore_color.rgb:
                        rgb = fill.fore_color.rgb
                        r, g, b = rgb[0], rgb[1], rgb[2]
                        hex_c = f"#{rgb}"
                        if calculate_luminance(r, g, b) < 130:
                            found_backgrounds.append(hex_c)
            except Exception:
                pass

            for shape in slide.shapes:
                # Check shape fills for dark backgrounds or accents
                try:
                    if shape.fill and hasattr(shape.fill, 'fore_color') and shape.fill.fore_color and shape.fill.fore_color.rgb:
                        rgb = shape.fill.fore_color.rgb
                        r, g, b = rgb[0], rgb[1], rgb[2]
                        hex_c = f"#{rgb}"
                        if calculate_luminance(r, g, b) < 130:
                            found_backgrounds.append(hex_c)
                        elif is_vibrant_accent(r, g, b):
                            found_accents.append(hex_c)
                except Exception:
                    pass

                # Check shape border / stroke lines for accents
                try:
                    if shape.line and shape.line.fill and shape.line.fill.fore_color and shape.line.fill.fore_color.rgb:
                        rgb = shape.line.fill.fore_color.rgb
                        r, g, b = rgb[0], rgb[1], rgb[2]
                        if is_vibrant_accent(r, g, b):
                            found_accents.append(f"#{rgb}")
                except Exception:
                    pass

                # Check text formatting for primary color, light text, and font family
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        # Slide section structure (first short line per slide), used by planning to follow the template's flow
                        txt = paragraph.text.strip()
                        if txt and len(txt) <= 70 and not txt.isdigit() and not title_found:
                            if txt not in detected_sections and not txt.startswith("By ") and len(txt) > 2:
                                detected_sections.append(txt)
                                title_found = True

                        if paragraph.font.name:
                            heading_font = paragraph.font.name
                        try:
                            if paragraph.font.color and paragraph.font.color.rgb:
                                rgb = paragraph.font.color.rgb
                                r, g, b = rgb[0], rgb[1], rgb[2]
                                hex_c = f"#{rgb}"
                                lum = calculate_luminance(r, g, b)
                                if is_vibrant_accent(r, g, b):
                                    found_accents.append(hex_c)
                                elif lum > 180:
                                    # White or very light text -> Template is a Dark Theme!
                                    has_light_text = True
                                elif lum < 100:
                                    primary_hex = hex_c
                        except Exception:
                            pass

        # Determine dominant background (dark vs light)
        if found_backgrounds:
            background_hex = found_backgrounds[0]
        elif has_light_text:
            # Light text strongly indicates a dark background presentation template
            background_hex = "#1E1E24"

        bg_clean = background_hex.lstrip('#')
        if len(bg_clean) == 6:
            r_bg, g_bg, b_bg = int(bg_clean[:2], 16), int(bg_clean[2:4], 16), int(bg_clean[4:6], 16)
            if calculate_luminance(r_bg, g_bg, b_bg) < 130 or has_light_text:
                # Dark mode template detected!
                if background_hex == "#FFFFFF":
                    background_hex = "#1E1E24"
                text_hex = "#FFFFFF"
                primary_hex = "#FFFFFF"
            else:
                text_hex = "#222222"

        if found_accents:
            accent_hex = found_accents[0]

        secondary_hex = "#8A94A6"
        body_font = heading_font

        # Prefer the brand colours and fonts declared in a customised theme; the slide scan above
        # remains the fallback for stock themes and still decides dark vs light backgrounds.
        theme = read_theme(file_path)
        if theme and not theme["is_stock"]:
            palette = palette_from_theme(theme)
            accent_hex = palette["accent_hex"]
            secondary_hex = palette["secondary_hex"]
            if text_hex != "#FFFFFF":  # light deck: theme colours are readable as-is
                primary_hex = palette["primary_hex"]
                background_hex = palette["background_hex"]
                text_hex = palette["text_hex"]
            heading_font = theme["heading_font"] or heading_font
            body_font = theme["body_font"] or body_font
        elif theme:
            body_font = theme["body_font"] or body_font

        raw_data: Dict[str, Any] = {
            "template_type": "PPTX",
            "source_file_name": base_name,
            "color_palette": {
                "primary_hex": primary_hex,
                "secondary_hex": secondary_hex,
                "accent_hex": accent_hex,
                "background_hex": background_hex,
                "text_hex": text_hex
            },
            "typography": {
                "heading_font": heading_font,
                "body_font": body_font,
                "heading_sizes": {"title": 28, "h1": 22, "h2": 18, "h3": 14, "body": 12}
            },
            "layout_rules": {
                "slide_layouts_available": available_layouts if available_layouts else ["Title Slide", "Title and Content"],
                "max_bullets_per_slide": 5,
                "table_header_shading_hex": primary_hex,
                "table_zebra_shading_hex": "#F4F7FA",
                "callout_box_style": "left_accent_border"
            },
            "detected_section_hierarchy": detected_sections
        }

        return normalize_guidance(raw_data)
    except Exception as e:
        print(f"Warning: PPTX parsing error on {file_path}: {e}. Returning default theme.")
        return normalize_guidance()

